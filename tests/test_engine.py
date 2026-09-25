"""End-to-end tests for the engine.

Several of these are regressions for specific defects found in the original
Grande. Where that is so, the test says which one and what the old behaviour was,
so the test explains itself when it fails in five years' time.
"""

from __future__ import annotations

import csv
import os
import sys
import time
from pathlib import Path

import openpyxl
import pytest

from grande.engine import export as export_engine
from grande.engine import pivot as pivot_engine
from grande.engine import query as query_engine
from grande.engine import ingest as ingest_module
from grande.engine import transform as transform_engine
from grande.engine.ingest import ingest, sniff
from grande.engine.session import Workspace
from grande.engine.sql import SqlError, ident

ROWS = [
    # item,     region, amount,     account, note,             big_id
    ["Sale",    "North", "2500.00", "07001", "ok",             "8931234567890123456"],
    ["Refund",  "North", "-1500.00", "00123", "=cmd|'/c calc'!A0", "8931234567890123457"],
    ["Sale",    "South", "1000.50", "40199", "+44 7700 900123", "8931234567890123458"],
    ["Refund",  "South", "-250.25", "00042", "needs review",   "8931234567890123459"],
    ["Sale",    "East",  "3000.00", "90210", "",               "8931234567890123460"],
]
HEADER = ["item", "region", "amount", "account", "note", "big_id"]


@pytest.fixture
def workspace(tmp_path):
    ws = Workspace(tmp_path / "ws")
    yield ws
    ws.close()


@pytest.fixture
def source(tmp_path):
    path = tmp_path / "ledger.csv"
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADER)
        writer.writerows(ROWS)
    return path


@pytest.fixture
def dataset(workspace, source):
    return ingest(workspace, source)


# ------------------------------------------------------------------- ingest


def test_sniff_detects_csv_shape(source):
    result = sniff(source)
    assert result.kind == "csv"
    assert result.delimiter == ","
    assert result.has_header is True
    assert [c["name"] for c in result.columns] == HEADER


def test_ingest_reads_every_row(dataset):
    assert dataset.row_count == len(ROWS)
    assert [c.name for c in dataset.columns] == HEADER


def test_leading_zeros_stay_text(dataset):
    """`07001` must not become 7001. Account numbers and ZIP codes depend on it."""
    account = dataset.column("account")
    assert account.kind == "text", "a zero-padded column must not be typed as a number"


def test_amount_is_numeric_despite_negatives(dataset):
    """Regression, original splitter.py:145.

    The original tested for a leading '-' (a formula character) *before* trying a
    numeric conversion, so negative numbers were written to Excel as text while
    positive ones were written as numbers. SUM then silently skipped the refunds.
    """
    assert dataset.column("amount").kind == "number"


# -------------------------------------------------------------------- query


def test_filter_and_count(workspace, dataset):
    filters = [{"column": "region", "op": "in", "values": ["North"]}]
    assert query_engine.count_filtered(workspace, dataset, filters) == 2


def test_summary_includes_negative_values(workspace, dataset):
    """The whole point: refunds must be part of the sum."""
    stats = query_engine.summary(workspace, dataset, "amount")
    assert stats["sum"] == pytest.approx(2500.00 - 1500.00 + 1000.50 - 250.25 + 3000.00)
    assert stats["min"] == pytest.approx(-1500.00)


def test_search_across_columns(workspace, dataset):
    filters = [{"op": "any_contains", "value": "refund", "columns": HEADER}]
    assert query_engine.count_filtered(workspace, dataset, filters) == 2


def test_unsorted_means_file_order(workspace, tmp_path):
    """Rows must come back in the order they appear in the file.

    Setting DuckDB's ``preserve_insertion_order = false`` is faster, but it
    lets rows come back in an unrelated order — so "row 1" would not be the
    first row of the user's file, and "keep the first duplicate" would keep an
    arbitrary one. Large enough here that the reader is genuinely parallel.
    """
    path = tmp_path / "ordered.csv"
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["n", "filler"])
        for i in range(120_000):
            writer.writerow([i, "x" * 20])
    ws = Workspace(tmp_path / "ws_order")
    try:
        ds = ingest(ws, path)
        page = query_engine.page(ws, ds, limit=50)
        assert [row[0] for row in page["rows"]] == list(range(50))
        tail = query_engine.page(ws, ds, offset=119_990, limit=10)
        assert [row[0] for row in tail["rows"]] == list(range(119_990, 120_000))
    finally:
        ws.close()


def test_sort_applies_to_whole_table(workspace, dataset):
    page = query_engine.page(
        workspace, dataset, limit=1, sort=[{"column": "amount", "direction": "asc"}]
    )
    assert page["rows"][0][HEADER.index("amount")] == pytest.approx(-1500.00)


def test_like_wildcards_are_escaped(workspace, dataset):
    """A literal % in a search box must not match everything."""
    filters = [{"column": "note", "op": "contains", "value": "%"}]
    assert query_engine.count_filtered(workspace, dataset, filters) == 0


# ---------------------------------------------------------------- injection


def test_identifier_quoting_doubles_quotes():
    """Regression, original pivot_engine.py:100, which used f'\"{col}\"' unescaped."""
    assert ident('a"b') == '"a""b"'
    with pytest.raises(SqlError):
        ident("")


def test_column_name_with_quote_is_safe(workspace, tmp_path):
    path = tmp_path / "odd.csv"
    path.write_text('id,we"ird\n1,x\n2,y\n', encoding="utf-8")
    ds = ingest(workspace, path)
    assert any('we"ird' in c.name for c in ds.columns)
    name = [c.name for c in ds.columns if c.name != "id"][0]
    # Would be a syntax error, or worse, without correct quoting.
    assert query_engine.count_filtered(
        workspace, ds, [{"column": name, "op": "is_not_null"}]
    ) == 2


def test_filter_value_cannot_inject(workspace, dataset):
    nasty = "'; DROP TABLE x; --"
    assert query_engine.count_filtered(
        workspace, dataset, [{"column": "note", "op": "contains", "value": nasty}]
    ) == 0
    assert query_engine.count_filtered(workspace, dataset, []) == len(ROWS)


def test_unknown_aggregator_is_rejected(workspace, dataset):
    with pytest.raises(pivot_engine.PivotError):
        pivot_engine.compute(
            workspace, dataset, rows=["region"],
            values=[{"column": "amount", "agg": "sum); DROP TABLE t; --"}],
        )


# -------------------------------------------------------------------- pivot


def test_pivot_groups_and_totals(workspace, dataset):
    result = pivot_engine.compute(
        workspace, dataset, rows=["region"],
        values=[{"column": "amount", "agg": "sum"}],
    )
    by_region = {row["keys"][0]: row["cells"][0] for row in result["rows"]}
    assert by_region["North"] == pytest.approx(1000.00)   # 2500 - 1500
    assert by_region["South"] == pytest.approx(750.25)    # 1000.50 - 250.25
    assert result["grand_total"]["cells"][0] == pytest.approx(4750.25)


def test_average_grand_total_is_not_an_average_of_averages(workspace, dataset):
    """Regression, original app.js:909-984.

    The original produced totals by summing the cells on screen. For AVG that
    gives the mean of the group means, which is only correct when every group is
    the same size. Here the total is computed from the underlying rows.
    """
    result = pivot_engine.compute(
        workspace, dataset, rows=["region"],
        values=[{"column": "amount", "agg": "avg"}],
    )
    true_mean = sum(float(r[2]) for r in ROWS) / len(ROWS)
    assert result["grand_total"]["cells"][0] == pytest.approx(true_mean)

    means = [row["cells"][0] for row in result["rows"]]
    assert result["grand_total"]["cells"][0] != pytest.approx(sum(means) / len(means))


def test_crosstab_row_total_matches_its_cells(workspace, dataset):
    result = pivot_engine.compute(
        workspace, dataset, rows=["region"], column="item",
        values=[{"column": "amount", "agg": "sum"}],
    )
    total_index = next(
        i for i, c in enumerate(result["columns"]) if c["role"] == "row_total"
    )
    cell_indexes = [i for i, c in enumerate(result["columns"]) if c["role"] == "cell"]
    for row in result["rows"]:
        cells = [row["cells"][i] or 0 for i in cell_indexes]
        assert row["cells"][total_index] == pytest.approx(sum(cells))


def test_percent_of_total_sums_to_one(workspace, dataset):
    result = pivot_engine.compute(
        workspace, dataset, rows=["region"],
        values=[{"column": "amount", "agg": "sum", "show_as": "percent_of_total"}],
    )
    assert sum(row["cells"][0] for row in result["rows"]) == pytest.approx(1.0)


def test_other_bucket_keeps_totals_whole(workspace, dataset):
    """Regression: the original dropped categories past the cap, so the row
    totals no longer matched the data. Here the remainder becomes 'Other'."""
    result = pivot_engine.compute(
        workspace, dataset, rows=["item"], column="region",
        values=[{"column": "amount", "agg": "sum"}],
        max_columns=1,
    )
    assert result["has_other"] is True
    total_index = next(i for i, c in enumerate(result["columns"]) if c["role"] == "row_total")
    kept = [i for i, c in enumerate(result["columns"]) if c["role"] in {"cell", "other"}]
    for row in result["rows"]:
        assert row["cells"][total_index] == pytest.approx(
            sum(row["cells"][i] or 0 for i in kept)
        )


def test_group_dates_by_month(workspace, tmp_path):
    """Regression for a verified gap: row fields were raw identifiers, so a date
    column grouped by individual day and "sales by month" was impossible."""
    path = tmp_path / "dated.csv"
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["when", "amount"])
        for day in (1, 5, 28):
            writer.writerow([f"2024-01-{day:02d}", 10])
        for day in (2, 17):
            writer.writerow([f"2024-02-{day:02d}", 100])
    ws = Workspace(tmp_path / "ws_dates")
    try:
        ds = ingest(ws, path)
        assert ds.column("when").kind == "date"

        ungrouped = pivot_engine.compute(
            ws, ds, rows=[{"column": "when"}], values=[{"column": "amount", "agg": "sum"}]
        )
        assert ungrouped["row_count"] == 5          # one row per distinct day

        monthly = pivot_engine.compute(
            ws, ds,
            rows=[{"column": "when", "group": "month"}],
            values=[{"column": "amount", "agg": "sum"}],
        )
        assert monthly["row_count"] == 2
        assert [r["cells"][0] for r in monthly["rows"]] == [30, 200]
        assert monthly["grand_total"]["cells"][0] == 230

        yearly = pivot_engine.compute(
            ws, ds,
            rows=[{"column": "when", "group": "year"}],
            values=[{"column": "amount", "agg": "sum"}],
        )
        assert yearly["row_count"] == 1
        assert yearly["rows"][0]["cells"][0] == 230
    finally:
        ws.close()


def test_group_numbers_into_bands(workspace, dataset):
    result = pivot_engine.compute(
        workspace, dataset,
        rows=[{"column": "amount", "bin": 1000}],
        values=[{"agg": "count_all"}],
    )
    assert sum(r["cells"][0] for r in result["rows"]) == len(ROWS)
    assert result["grand_total"]["cells"][0] == len(ROWS)


def test_two_column_fields_nest(workspace, dataset):
    """Excel allows several column fields; headers stack and totals still hold."""
    result = pivot_engine.compute(
        workspace, dataset,
        rows=["region"], columns=["item", "status" if dataset.column("status") else "item"],
        values=[{"column": "amount", "agg": "sum"}],
    )
    total_index = next(i for i, c in enumerate(result["columns"]) if c["role"] == "row_total")
    cells = [i for i, c in enumerate(result["columns"]) if c["role"] in {"cell", "other"}]
    for row in result["rows"]:
        assert row["cells"][total_index] == pytest.approx(
            sum(row["cells"][i] or 0 for i in cells)
        )


def test_grouped_drill_down_uses_a_range(workspace, tmp_path):
    """A month cell must select the whole month, not rows equal to the 1st."""
    path = tmp_path / "d2.csv"
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["when", "amount"])
        for day in range(1, 29):
            writer.writerow([f"2024-03-{day:02d}", 1])
        writer.writerow(["2024-04-02", 1])
    ws = Workspace(tmp_path / "ws_drill")
    try:
        ds = ingest(ws, path)
        data = pivot_engine.drill_down(
            ws, ds,
            row_fields=[{"column": "when", "group": "month"}],
            keys=["2024-03-01"],
        )
        assert data["filtered_rows"] == 28
    finally:
        ws.close()


def test_drill_down_returns_underlying_rows(workspace, dataset):
    data = pivot_engine.drill_down(
        workspace, dataset, row_fields=["region"], keys=["North"]
    )
    assert data["filtered_rows"] == 2


# ---------------------------------------------------------------- transform


def test_remove_duplicates(workspace, dataset):
    before = dataset.row_count
    transform_engine.remove_duplicates(workspace, dataset, columns=["region"])
    assert dataset.row_count == 3            # North, South, East
    transform_engine.undo(workspace, dataset)
    assert dataset.row_count == before


def test_find_replace_and_undo(workspace, dataset):
    transform_engine.find_replace(
        workspace, dataset, columns=["item"], find="Sale", replace="Order"
    )
    page = query_engine.page(workspace, dataset, limit=10)
    items = {row[0] for row in page["rows"]}
    assert "Order" in items and "Sale" not in items
    transform_engine.undo(workspace, dataset)
    page = query_engine.page(workspace, dataset, limit=10)
    assert "Sale" in {row[0] for row in page["rows"]}


def test_change_type_reports_unconvertible(workspace, dataset):
    transform_engine.change_type(workspace, dataset, column="note", to="number")
    step = transform_engine.recipe(dataset)[-1]
    assert step["params"]["unconverted"] > 0


# ------------------------------------------------------------------- export


def _read(path):
    book = openpyxl.load_workbook(path)
    sheet = book.active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


@pytest.mark.parametrize("fast", [False, True])
def test_excel_export_types(workspace, dataset, tmp_path, fast):
    if fast and not export_engine.excel_writer_available():
        pytest.skip("DuckDB Excel writer not available offline")

    result = export_engine.export_excel(
        workspace, dataset, directory=str(tmp_path / "out"),
        base_name="ledger", sort=[{"column": "item", "direction": "asc"}],
        fast_writer=fast,
    )
    assert result["status"] == "ok"
    assert len(result["files"]) == 1

    rows = _read(result["files"][0]["path"])
    assert rows[0] == HEADER

    amounts = [r[HEADER.index("amount")] for r in rows[1:]]
    assert all(isinstance(v, (int, float)) for v in amounts), (
        "every amount must be a number — the original wrote negatives as text"
    )
    assert sum(amounts) == pytest.approx(4750.25)

    accounts = [r[HEADER.index("account")] for r in rows[1:]]
    assert all(isinstance(v, str) for v in accounts)
    assert "07001" in accounts, "leading zero lost"

    # An 19-digit id is a perfectly ordinary BIGINT to DuckDB, but Excel keeps
    # only 15 significant digits, so writing it as a number silently rewrites
    # the tail as zeros and shows 8.93123E+18.
    ids = [r[HEADER.index("big_id")] for r in rows[1:]]
    assert all(isinstance(v, str) for v in ids), (
        "long identifiers must be written as text or Excel rounds them"
    )
    assert "8931234567890123456" in ids
    assert any("15 digits" in w for w in result["warnings"])


def test_excel_export_never_executes_formulas(workspace, dataset, tmp_path):
    result = export_engine.export_excel(
        workspace, dataset, directory=str(tmp_path / "out"),
        base_name="ledger", fast_writer=False,
    )
    book = openpyxl.load_workbook(result["files"][0]["path"])
    sheet = book.active
    column = HEADER.index("note") + 1
    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row=row, column=column)
        assert cell.data_type != "f", "a cell was written as a formula"


def test_multipart_split_covers_every_row(workspace, dataset, tmp_path):
    result = export_engine.export_excel(
        workspace, dataset, directory=str(tmp_path / "parts"),
        base_name="ledger", rows_per_file=2, suffix_style="alpha",
        fast_writer=False,
    )
    assert [f["name"] for f in result["files"]] == [
        "ledger_a.xlsx", "ledger_b.xlsx", "ledger_c.xlsx",
    ]
    assert sum(f["rows"] for f in result["files"]) == len(ROWS)
    for file in result["files"]:
        rows = _read(file["path"])
        assert rows[0] == HEADER, "every part must repeat the header"


def test_export_respects_filters_and_columns(workspace, dataset, tmp_path):
    result = export_engine.export_excel(
        workspace, dataset, directory=str(tmp_path / "filtered"),
        base_name="north", columns=["item", "amount"],
        filters=[{"column": "region", "op": "in", "values": ["North"]}],
        fast_writer=False,
    )
    rows = _read(result["files"][0]["path"])
    assert rows[0] == ["item", "amount"]
    assert len(rows) == 3          # header + 2 North rows


def test_rows_per_file_capped_at_excel_limit(workspace, dataset, tmp_path):
    plan = export_engine.plan_export(
        workspace, dataset, directory=str(tmp_path),
        rows_per_file=5_000_000, extension=".xlsx",
    )
    assert plan.rows_per_file == 1_048_575
    assert any("Excel holds at most" in w for w in plan.warnings)


def test_single_part_has_no_suffix(workspace, dataset, tmp_path):
    plan = export_engine.plan_export(
        workspace, dataset, directory=str(tmp_path), base_name="one", extension=".xlsx"
    )
    assert plan.filenames == ["one.xlsx"]


def test_flat_exports(workspace, dataset, tmp_path):
    for fmt, suffix in (("csv", ".csv"), ("parquet", ".parquet"), ("json", ".jsonl")):
        result = export_engine.export_flat(
            workspace, dataset, directory=str(tmp_path / fmt), base_name="x", fmt=fmt
        )
        assert result["status"] == "ok"
        assert Path(result["files"][0]["path"]).suffix == suffix
        assert Path(result["files"][0]["path"]).exists()


def test_long_cells_truncated_and_reported(workspace, tmp_path):
    path = tmp_path / "long.csv"
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["id", "blob"])
        writer.writerow(["1", "x" * 40_000])
    ws = Workspace(tmp_path / "ws2")
    try:
        ds = ingest(ws, path)
        result = export_engine.export_excel(
            ws, ds, directory=str(tmp_path / "o"), base_name="long", fast_writer=False
        )
        assert any("32,767" in w for w in result["warnings"]), (
            "silent truncation is what the original did; it must be reported"
        )
    finally:
        ws.close()


# ------------------------------------------------------------------ formats


def test_excel_input_roundtrip(workspace, dataset, tmp_path):
    """Excel in, Excel out — the original could not read .xlsx at all."""
    exported = export_engine.export_excel(
        workspace, dataset, directory=str(tmp_path / "rt"),
        base_name="rt", fast_writer=False,
    )
    reopened = ingest(workspace, exported["files"][0]["path"])
    assert reopened.row_count == len(ROWS)
    assert [c.name for c in reopened.columns] == HEADER
    assert reopened.column("account").kind == "text", "leading zeros lost on reload"
    # Regression: decimal columns came back rounded to whole numbers.
    amounts = [r[0] for r in workspace.execute(f"SELECT amount FROM {reopened.relation} ORDER BY rowid")]
    assert amounts == [float(r[2]) for r in ROWS]


def test_semicolon_file_is_detected(workspace, tmp_path):
    path = tmp_path / "euro.csv"
    path.write_text("a;b;c\n1;2;3\n4;5;6\n", encoding="utf-8")
    result = sniff(path)
    assert result.delimiter == ";"
    ds = ingest(workspace, path, delimiter=";")
    assert ds.row_count == 2 and len(ds.columns) == 3


def test_big_integers_survive_the_json_boundary(workspace, dataset):
    """JavaScript rounds past 2^53, so large ints go to the browser as strings."""
    page = query_engine.page(workspace, dataset, limit=1, columns=["big_id"])
    value = page["rows"][0][0]
    assert isinstance(value, str)
    assert value.startswith("89312345678901234")


def _start_holder(directory, tmp_path):
    """Open `directory` in another process and wait until it holds the lock."""
    import subprocess

    ready = tmp_path / f"ready-{Path(directory).name}"
    child = subprocess.Popen(
        [sys.executable, str(Path(__file__).parent / "_hold_workspace.py"),
         str(directory), str(ready)],
        stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        cwd=str(Path(__file__).parent.parent),
    )
    for _ in range(120):
        if ready.exists():
            return child
        if child.poll() is not None:
            pytest.skip("helper process could not open a workspace")
        time.sleep(0.1)
    child.kill()
    pytest.skip("helper process did not start in time")


def test_second_instance_gets_its_own_workspace(tmp_path, monkeypatch):
    """Launching Grande twice must not fail.

    DuckDB's file lock is held per process, so a second *process* cannot share
    the default cache. Before this, a second double-click ended in a DuckDB
    traceback; now the second run quietly gets a cache of its own.
    """
    from grande.engine import session as session_module

    default = tmp_path / "cache"
    child = _start_holder(default, tmp_path)
    try:
        monkeypatch.setattr(session_module, "_default_cache_dir", lambda: default)
        second = Workspace()
        try:
            assert second._temporary is True, "should have fallen back to its own cache"
            assert second.dir != default
            assert second.execute_one("SELECT 42")[0] == 42     # genuinely usable
        finally:
            second.close()
        # A per-run cache is removed with the run that created it.
        assert not second.dir.exists()
    finally:
        child.kill()
        child.wait(timeout=15)


def test_explicit_workspace_in_use_is_reported(tmp_path):
    """An explicit --workspace must say it is busy rather than silently moving."""
    from grande.engine.session import WorkspaceBusy

    shared = tmp_path / "shared"
    child = _start_holder(shared, tmp_path)
    try:
        with pytest.raises(WorkspaceBusy):
            Workspace(shared)
    finally:
        child.kill()
        child.wait(timeout=15)


# ------------------------------------------------------------- type changing


def _typed(tmp_path, name, values, header="value"):
    path = tmp_path / f"{name}.csv"
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["id", header])
        for i, v in enumerate(values):
            writer.writerow([i, v])
    return path


def test_non_iso_dates_convert(tmp_path):
    """Regression: a plain TRY_CAST only understands ISO-ish text.

    `15/03/2024`, `15-Mar-2024`, `Mar 15, 2024`, `15.03.2024` and `20240315` all
    became empty, so a perfectly ordinary export lost most of its date column
    and was told only afterwards, with no indication of which values.
    """
    values = ["2024-03-15", "15/03/2024", "15-Mar-2024", "Mar 15, 2024",
              "15.03.2024", "20240315", "2024/03/15"]
    ws = Workspace(tmp_path / "ws_dates2")
    try:
        ds = ingest(ws, _typed(tmp_path, "d", values), )
        preview = transform_engine.preview_change_type(ws, ds, column="value", to="date")
        assert preview["failed"] == 0, f"still failing: {preview['failures']}"
        assert preview["converted"] == len(values)

        transform_engine.change_type(ws, ds, column="value", to="date")
        page = query_engine.page(ws, ds, limit=20, columns=["value"])
        assert {str(r[0]) for r in page["rows"]} == {"2024-03-15"}
    finally:
        ws.close()


def test_unconvertible_values_are_named_before_applying(tmp_path):
    ws = Workspace(tmp_path / "ws_named")
    try:
        ds = ingest(ws, _typed(tmp_path, "n", ["2024-03-15", "N/A", "N/A", "TBC"]))
        preview = transform_engine.preview_change_type(ws, ds, column="value", to="date")
        assert preview["failed"] == 3
        found = {f["value"]: f["count"] for f in preview["failures"]}
        assert found == {"N/A": 2, "TBC": 1}
        # Nothing has changed yet: a preview must not touch the data.
        assert ds.column("value").kind == "text"
    finally:
        ws.close()


def test_day_month_ambiguity_is_reported(tmp_path):
    """03/05/2024 is 3 May or 5 March. Guessing silently is not acceptable.

    Loaded as text (as happens whenever mixed layouts defeat the sniffer), the
    preview must say that the reading is ambiguous before anything is changed.
    """
    ws = Workspace(tmp_path / "ws_ambig")
    try:
        ds = ingest(ws, _typed(tmp_path, "a", ["03/05/2024", "07/08/2024"]),
                    all_varchar=True)
        preview = transform_engine.preview_change_type(ws, ds, column="value", to="date")
        assert preview["ambiguous"] is True
        assert preview["ambiguous_rows"] == 2

        transform_engine.change_type(ws, ds, column="value", to="date", day_first=True)
        first = query_engine.page(ws, ds, limit=1, columns=["value"])["rows"][0][0]
        assert str(first) == "2024-05-03"          # 3 May
    finally:
        ws.close()


def test_month_first_reading(tmp_path):
    ws = Workspace(tmp_path / "ws_mf")
    try:
        ds = ingest(ws, _typed(tmp_path, "m", ["03/05/2024"]), all_varchar=True)
        transform_engine.change_type(ws, ds, column="value", to="date", day_first=False)
        first = query_engine.page(ws, ds, limit=1, columns=["value"])["rows"][0][0]
        assert str(first) == "2024-03-05"          # 5 March
    finally:
        ws.close()


def test_loader_says_which_date_reading_it_used(tmp_path):
    """DuckDB's CSV sniffer settles on a date layout silently.

    For 03/05/2024 that choice moves every date by two months, so the load has
    to say which way it read them, and let the file be reopened the other way.
    """
    from grande.engine.ingest import detect_date_format

    path = _typed(tmp_path, "warn", ["03/05/2024", "07/08/2024"], header="when")
    assert detect_date_format(path) in ingest_module.AMBIGUOUS_DATE_FORMATS

    probe = sniff(path)
    assert probe.date_format is not None
    assert any("day first" in n or "month first" in n for n in probe.notes)

    ws = Workspace(tmp_path / "ws_warn")
    try:
        guessed = ingest(ws, path)
        assert guessed.column("when").kind == "date"
        assert any("read day first" in w or "read month first" in w
                   for w in guessed.warnings), guessed.warnings

        # And the choice can be overridden when reopening.
        forced = ingest(ws, path, date_format="%m/%d/%Y")
        first = query_engine.page(ws, forced, limit=1, columns=["when"])["rows"][0][0]
        assert str(first) == "2024-03-05"
    finally:
        ws.close()


def test_keep_original_preserves_unconvertible_values(tmp_path):
    ws = Workspace(tmp_path / "ws_keep")
    try:
        ds = ingest(ws, _typed(tmp_path, "k", ["2024-03-15", "not known"]))
        transform_engine.change_type(ws, ds, column="value", to="date", keep_original=True)
        names = [c.name for c in ds.columns]
        assert "value (original)" in names
        page = query_engine.page(ws, ds, limit=10, columns=["value (original)"])
        assert "not known" in {r[0] for r in page["rows"]}, "the value must survive"
    finally:
        ws.close()


def test_messy_numbers_convert(tmp_path):
    """Spreadsheet exports carry separators, currency marks and (1,234) negatives."""
    ws = Workspace(tmp_path / "ws_num")
    try:
        ds = ingest(ws, _typed(tmp_path, "num", ["1,234.50", "$2,000", "(1,500.00)", "42"]))
        preview = transform_engine.preview_change_type(ws, ds, column="value", to="number")
        assert preview["failed"] == 0, f"still failing: {preview['failures']}"
        transform_engine.change_type(ws, ds, column="value", to="number")
        values = [r[0] for r in query_engine.page(ws, ds, limit=10, columns=["value"])["rows"]]
        assert values == pytest.approx([1234.50, 2000.0, -1500.0, 42.0])
    finally:
        ws.close()


# ------------------------------------------------------------- the SQL console


def test_sql_console_allows_reads():
    from grande.engine.sql import check_read_only

    for statement in ("SELECT * FROM data",
                      "  with x as (select 1) select * from x  ",
                      "DESCRIBE data",
                      "SUMMARIZE data",
                      "SELECT * FROM data; "):
        assert check_read_only(statement)


def test_sql_console_refuses_writes():
    from grande.engine.sql import check_read_only

    for statement in ("COPY data TO 'x.csv'",
                      "INSTALL httpfs",
                      "ATTACH 'other.db'",
                      "DROP TABLE data",
                      "CREATE TABLE t AS SELECT 1",
                      "SET disabled_filesystems = ''",
                      "SELECT 1; DROP TABLE data"):
        with pytest.raises(SqlError):
            check_read_only(statement)


def test_a_value_that_looks_like_a_keyword_is_fine():
    """The check reads the statement, so a literal must not trip it."""
    from grande.engine.sql import check_read_only

    assert check_read_only("SELECT * FROM data WHERE note = 'please delete this'")
    assert check_read_only("SELECT * FROM data -- drop everything")


def test_running_sql_does_not_break_opening_files(tmp_path):
    """Regression: the console used to sandbox itself with

        SET disabled_filesystems = 'LocalFileSystem'

    which is global to the DuckDB instance. One query and every later file open
    failed with a permission error until Grande was restarted.
    """
    ws = Workspace(tmp_path / "ws_sqlfs")
    try:
        first = ingest(ws, _typed(tmp_path, "one", ["a", "b"]))
        assert first.row_count == 2

        cur = ws.cursor()
        try:
            cur.execute(f"SELECT count(*) FROM {first.relation}").fetchone()
        finally:
            cur.close()

        second = ingest(ws, _typed(tmp_path, "two", ["c", "d"]))
        assert second.row_count == 2, "a file must still be openable after a query"
    finally:
        ws.close()


# ------------------------------------------------------ review regressions


def test_deleting_matches_keeps_rows_the_filter_did_not_match(tmp_path):
    """Regression: NOT (rev > 100) is NULL for a blank rev, so 'Delete matching
    rows' deleted every blank row along with the matches."""
    ws = Workspace(tmp_path / "ws_invert")
    try:
        path = tmp_path / "rev.csv"
        path.write_text("id,rev\n1,50\n2,500\n3,\n4,20\n", encoding="utf-8")
        ds = ingest(ws, path)
        rule = [{"column": "rev", "op": "gt", "value": 100}]
        transform_engine.keep_filtered(ws, ds, filters=rule, invert=True)
        ids = [r[0] for r in query_engine.page(ws, ds, limit=10, columns=["id"])["rows"]]
        assert sorted(ids) == [1, 3, 4]

        transform_engine.undo(ws, ds)
        transform_engine.keep_filtered(ws, ds, filters=rule)
        ids = [r[0] for r in query_engine.page(ws, ds, limit=10, columns=["id"])["rows"]]
        assert ids == [2]
    finally:
        ws.close()


def test_decimal_commas_are_read_as_decimals(tmp_path):
    """Regression: every comma was stripped, so 12,5 became 125 and 1.234,56
    became 1.23456 — counted as converted — and 'SKU-123' became -123."""
    ws = Workspace(tmp_path / "ws_eu")
    try:
        values = ["12,5", "1.234,56", "0,75", "(1.234,56)", "1.234", "1,234", "€ 3,5"]
        ds = ingest(ws, _typed(tmp_path, "eu", values + ["SKU-123"]))
        preview = transform_engine.preview_change_type(ws, ds, column="value", to="number")
        assert [f["value"] for f in preview["failures"]] == ["SKU-123"]
        transform_engine.change_type(ws, ds, column="value", to="number")
        got = [r[0] for r in query_engine.page(ws, ds, limit=10, columns=["value"])["rows"]]
        assert got[:-1] == pytest.approx([12.5, 1234.56, 0.75, -1234.56, 1.234, 1234.0, 3.5])
        assert got[-1] is None
    finally:
        ws.close()


def test_export_never_replaces_the_source_unasked(workspace, dataset, source):
    """Regression: the form proposes the source's folder and name, and a
    same-format export silently wrote over the user's original file."""
    original = source.read_bytes()
    plan = export_engine.plan_export(
        workspace, dataset, directory=str(source.parent), base_name="ledger", extension=".csv",
    )
    assert plan.existing == ["ledger.csv"]
    assert plan.replaces_source == "ledger.csv"
    assert "would replace it" in plan.as_dict()["warnings"][0]

    with pytest.raises(export_engine.ExportError, match="file you opened"):
        export_engine.export_flat(
            workspace, dataset, directory=str(source.parent), base_name="ledger", fmt="csv",
            filters=[{"column": "item", "op": "in", "values": ["Sale"]}],
        )
    assert source.read_bytes() == original

    result = export_engine.export_flat(
        workspace, dataset, directory=str(source.parent), base_name="ledger", fmt="csv",
        overwrite=True,
    )
    assert source.read_bytes() != original
    assert result["warnings"][0] == "Replaced ledger.csv, the file you opened."


@pytest.mark.parametrize("fast", [False, True])
def test_excel_export_asks_before_replacing(workspace, dataset, tmp_path, fast):
    if fast and not export_engine.excel_writer_available():
        pytest.skip("DuckDB excel extension not available")
    out = tmp_path / "out"
    export_engine.export_excel(
        workspace, dataset, directory=str(out), base_name="book", fast_writer=fast,
    )
    before = (out / "book.xlsx").stat().st_mtime_ns
    with pytest.raises(export_engine.ExportError, match="already exists"):
        export_engine.export_excel(
            workspace, dataset, directory=str(out), base_name="book", fast_writer=fast,
        )
    assert (out / "book.xlsx").stat().st_mtime_ns == before
    result = export_engine.export_excel(
        workspace, dataset, directory=str(out), base_name="book", fast_writer=fast,
        overwrite=True,
    )
    assert result["status"] == "ok"


@pytest.mark.parametrize("fast", [False, True])
def test_excel_export_into_a_bracketed_folder(workspace, dataset, tmp_path, fast):
    """The staging file is read back by path, and DuckDB globs paths."""
    if fast and not export_engine.excel_writer_available():
        pytest.skip("DuckDB excel extension not available")
    out = tmp_path / "exports [2024]"
    result = export_engine.export_excel(
        workspace, dataset, directory=str(out), base_name="book", fast_writer=fast,
    )
    assert result["row_count"] == len(ROWS)
    assert (out / "book.xlsx").exists()
